"""Builds the student's lecture deck used by the engineer film: real Mechanics of Materials
content (Euler-Bernoulli beam bending), every number hand-checked. Output: deck.pptx.

Worked numbers (so they can be re-verified):
  Case A, simply supported, midspan point load P = 12 kN, L = 3.0 m, E = 200 GPa,
          I = 8.0e-6 m^4, c = 0.10 m
    M_max = PL/4 = 9.0 kN*m           sigma_max = M c / I = 112.5 MPa
    delta_max = P L^3 / (48 E I) = 12000*27/(48*200e9*8e-6) = 4.22 mm
  Case B, same beam as a cantilever with the load at the tip
    M_max = PL = 36 kN*m              sigma_max = 450 MPa
    delta_max = P L^3 / (3 E I) = 12000*27/(3*200e9*8e-6) = 67.5 mm
  Ratios: moment 4x, deflection 16x.
  Allowables: delta <= L/360 = 8.33 mm; sigma <= 250/1.5 = 166.7 MPa.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1B, 0x2A, 0x49)
BLUE = RGBColor(0x2F, 0x4D, 0x8C)
GREY = RGBColor(0x55, 0x5B, 0x66)
RED = RGBColor(0xB0, 0x32, 0x11)
LIGHT = RGBColor(0xEE, 0xF1, 0xF7)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=20, color=NAVY, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(6)
        if align:
            p.alignment = align
    return box


def slide_with_title(title, subtitle=None):
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    text(s, 0.7, 0.45, 12, 0.9, [title], size=34, bold=True)
    if subtitle:
        text(s, 0.7, 1.25, 12, 0.5, [subtitle], size=16, color=GREY)
    return s


def beam(slide, x, y, w, kind):
    """A simple beam schematic: a bar plus supports (kind: 'ss' or 'cantilever')."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.22))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    if kind == "ss":
        for sx in (x, x + w):
            tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(sx - 0.2), Inches(y + 0.22), Inches(0.4), Inches(0.36))
            tri.fill.solid()
            tri.fill.fore_color.rgb = GREY
            tri.line.fill.background()
        ax = x + w / 2
    else:
        wall = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - 0.25), Inches(y - 0.5), Inches(0.25), Inches(1.22))
        wall.fill.solid()
        wall.fill.fore_color.rgb = GREY
        wall.line.fill.background()
        ax = x + w
    arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(ax - 0.18), Inches(y - 0.95), Inches(0.36), Inches(0.9))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RED
    arrow.line.fill.background()
    text(slide, ax - 0.4, y - 1.35, 0.8, 0.4, ["P"], size=20, color=RED, bold=True, align=PP_ALIGN.CENTER)
    text(slide, x, y + 0.7, w, 0.4, ["L"], size=18, color=GREY, align=PP_ALIGN.CENTER)


# 1 — title
s = prs.slides.add_slide(blank)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
text(s, 0.9, 2.3, 11.5, 1.2, ["ME 3310 — Mechanics of Materials"], size=44, color=RGBColor(255, 255, 255), bold=True)
text(s, 0.9, 3.5, 11.5, 0.8, ["Lecture 9: Beam Bending & Deflection — my notes"], size=28, color=RGBColor(0xC9, 0xD4, 0xEE))
text(s, 0.9, 5.4, 11.5, 0.6, ["Marcus Reyes · Mechanical Engineering, Year 3"], size=18, color=RGBColor(0xC9, 0xD4, 0xEE))

# 2 — assumptions
s = slide_with_title("Euler–Bernoulli beam theory", "What the formulas on the next slides assume")
text(s, 0.7, 1.9, 11.8, 4.8, [
    "• Plane sections stay plane and perpendicular to the neutral axis",
    "• Small deflections and slopes; linear-elastic material (Hooke's law)",
    "• Governing equation:  E I · d²v/dx² = M(x)",
    "• Integrate twice, then fix the constants with boundary conditions:",
    "     – pinned / roller support:  v = 0",
    "     – fixed (built-in) support:  v = 0  AND  dv/dx = 0",
    "     – symmetry about midspan:  dv/dx = 0 at x = L/2",
], size=22)

# 3 — bending stress
s = slide_with_title("Bending stress", "Normal stress from the internal bending moment")
text(s, 0.7, 1.9, 11.8, 4.8, [
    "σ(y) = − M y / I",
    "Maximum at the outer fibre, y = c:     σ_max = M c / I = M / S",
    "Section modulus:  S = I / c",
    "Tension on one side of the neutral axis, compression on the other; σ = 0 on the axis.",
    "Design check:  σ_max ≤ σ_allow = σ_yield / FS",
], size=24)

# 4 — case A
s = slide_with_title("Case A — simply supported, point load at midspan")
beam(s, 1.2, 3.4, 5.6, "ss")
text(s, 7.6, 1.9, 5.2, 4.8, [
    "Reactions:  R_A = R_B = P / 2",
    "M_max = P L / 4   (at midspan)",
    "δ_max = P L³ / (48 E I)   (at x = L/2)",
    "Slope at the supports:  θ = P L² / (16 E I)",
], size=22)

# 5 — case B
s = slide_with_title("Case B — cantilever, point load at the tip")
beam(s, 1.6, 3.4, 5.2, "cantilever")
text(s, 7.6, 1.9, 5.2, 4.8, [
    "Reaction moment at the wall:  M_wall = P L",
    "M_max = P L   (at the fixed end)",
    "δ_max = P L³ / (3 E I)   (at the tip)",
    "Slope at the tip:  θ = P L² / (2 E I)",
], size=22)

# 6 — worked example A
s = slide_with_title("Worked example — Case A", "Steel W-section:  L = 3.0 m,  P = 12 kN,  E = 200 GPa,  I = 8.0×10⁻⁶ m⁴,  c = 0.10 m")
text(s, 0.7, 2.0, 11.8, 4.8, [
    "M_max = P L / 4 = (12 000 N)(3.0 m) / 4 = 9.0 kN·m",
    "σ_max = M c / I = (9 000)(0.10) / (8.0×10⁻⁶) = 112.5 MPa",
    "δ_max = P L³ / (48 E I) = (12 000)(27) / (48 · 200×10⁹ · 8.0×10⁻⁶) = 4.22 mm",
    "Checks:  δ ≤ L/360 = 8.33 mm  ✓      σ ≤ 250 / 1.5 = 166.7 MPa  ✓",
], size=22)

# 7 — worked example B
s = slide_with_title("Worked example — Case B (same beam, cantilever)", "Same L, P, E, I, c")
text(s, 0.7, 2.0, 11.8, 4.8, [
    "M_max = P L = 36 kN·m",
    "σ_max = M c / I = (36 000)(0.10) / (8.0×10⁻⁶) = 450 MPa   ✗  (yield is only 250 MPa)",
    "δ_max = P L³ / (3 E I) = (12 000)(27) / (3 · 200×10⁹ · 8.0×10⁻⁶) = 67.5 mm   ✗  (limit 8.33 mm)",
    "Cantilever vs simply supported:   moment 4× larger,   deflection 16× larger",
], size=22)

# 8 — open question
s = slide_with_title("Open question (mine)")
text(s, 0.7, 1.9, 11.8, 4.8, [
    "Why is the cantilever exactly 16× worse in deflection but only 4× worse in moment?",
    "I can recite both formulas — I want to be able to explain the factor from the",
    "boundary conditions, not from the table.",
], size=26, color=RED)

# 9 — homework
s = slide_with_title("Homework 9 — Problem 4 (not started)")
text(s, 0.7, 1.9, 11.8, 4.8, [
    "A simply supported steel beam, L = 4.5 m, carries a central point load P.",
    "E = 200 GPa,  I = 12×10⁻⁶ m⁴,  c = 0.12 m,  σ_yield = 250 MPa,  FS = 1.5.",
    "Deflection limit: δ_max ≤ L / 360.",
    "Find the largest P the beam can carry, and state which check governs.",
], size=24)

prs.save("deck.pptx")
print("wrote deck.pptx", len(prs.slides._sldIdLst), "slides")
