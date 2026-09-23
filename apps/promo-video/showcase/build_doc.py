"""Builds the student's lecture handout used by the showcase film: real Calculus II content on
integration by parts, every result hand-checked. Output: handout.docx.

  integral u dv = u v - integral v du
  x e^x dx        u = x,  dv = e^x dx     -> x e^x - e^x + C
  x cos x dx      u = x,  dv = cos x dx   -> x sin x + cos x + C
  ln x dx         u = ln x, dv = dx       -> x ln x - x + C
  HW: x^2 e^x dx  -> e^x (x^2 - 2x + 2) + C ;  x ln x dx -> (x^2/2) ln x - x^2/4 + C
"""
from docx import Document
from docx.shared import Pt, RGBColor

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(12)

h = doc.add_heading("MATH 1220 — Calculus II, Lecture 12: Integration by Parts", level=1)
doc.add_paragraph("Prof. Alvarez · my notes · Week 12")

doc.add_heading("The rule", level=2)
doc.add_paragraph("∫ u dv = u v − ∫ v du")
doc.add_paragraph(
    "Use it when the integrand is a product of two different kinds of function and a "
    "substitution won't untangle it. Split the integrand into u (which we differentiate) and "
    "dv (which we integrate), get du and v, then plug in."
)

doc.add_heading("Choosing u: LIATE", level=2)
doc.add_paragraph(
    "Pick u as whichever comes first: Logarithmic, Inverse trig, Algebraic, Trig, Exponential. "
    "Everything left over is dv."
)

doc.add_heading("Worked examples", level=2)
doc.add_paragraph("Example 1.  ∫ x e^x dx")
doc.add_paragraph("u = x, dv = e^x dx, so du = dx and v = e^x.")
doc.add_paragraph("∫ x e^x dx = x e^x − ∫ e^x dx = x e^x − e^x + C = e^x (x − 1) + C")
doc.add_paragraph("Example 2.  ∫ x cos x dx")
doc.add_paragraph("u = x, dv = cos x dx, so du = dx and v = sin x.")
doc.add_paragraph("∫ x cos x dx = x sin x − ∫ sin x dx = x sin x + cos x + C")
doc.add_paragraph("Example 3.  ∫ ln x dx")
doc.add_paragraph("u = ln x, dv = dx, so du = dx / x and v = x.")
doc.add_paragraph("∫ ln x dx = x ln x − ∫ 1 dx = x ln x − x + C")

doc.add_heading("Homework 12", level=2)
doc.add_paragraph("1.  ∫ x² e^x dx   (hint: apply the rule twice)")
doc.add_paragraph("2.  ∫ x ln x dx")
doc.add_paragraph("Midterm 3 is Thursday of week 13 — covers substitution, parts, and partial fractions.")

doc.save("handout.docx")
print("wrote handout.docx")
