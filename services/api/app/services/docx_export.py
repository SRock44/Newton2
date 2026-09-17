"""Builds a real Word `.docx` from a Newton note or document, so work a student did in
this app can be handed in, emailed to a professor, or opened on a school computer that
has Word and nothing else.

DESIGN DECISION -- this maps Markdown structure onto real Word styles rather than
dumping the text one-line-per-paragraph, because Newton's notes and documents genuinely
ARE Markdown, not incidental plain text:

  * app/services/documents.py's create_note() stores every Notepad note with
    mime_type="text/markdown", and apps/desktop/src/NotepadWindow.tsx edits it in a
    Write/Preview editor whose Preview tab is the app's Markdown renderer.
  * Both the Notepad preview and the Documents panel's detail pane render document
    content through <MessageContent>, i.e. ReactMarkdown + remark-gfm -- so headings,
    bold, and lists are exactly what the student already SEES.
  * The Notepad's highlight-to-act actions (Explain/Define/Summarize) insert model
    output straight into the note, and that output is Markdown by construction.

Flattening all of that would mean a student who wrote "## Causes" gets a literal "##
Causes" in their Word document -- visibly worse than what they were looking at a second
earlier. So `#`/`##`/`###` become real Word heading styles, `-`/`1.` become real Word
list styles, `**bold**`/`*italic*`/`` `code` `` become real runs, GFM tables become real
Word tables, and fenced code becomes monospace.

Deliberately NOT a full CommonMark implementation -- there is no nesting, no reference
links, no footnotes, no inline HTML. This is a line-oriented mapping of the constructs
that actually show up in student notes, and anything it doesn't recognize falls through
as ordinary paragraph text, which is always readable. A real Markdown parser would be a
new dependency and a much larger surface for what is, in the end, a convenience export.

Pure `title + content in, bytes out`: no DB session, no I/O, no model call.
"""

from __future__ import annotations

import io
import re

import docx
from docx.shared import Pt

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_QUOTE = re.compile(r"^\s*>\s?(.*)$")
_RULE = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")
_FENCE = re.compile(r"^\s*```")
_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_TABLE_DIVIDER = re.compile(r"^\s*\|[\s:|-]*-[\s:|-]*\|\s*$")

# One pass over a line's inline markup. Bold is matched before italic so "**x**" never
# gets read as an empty italic wrapping "*x*", and links are flattened to "text (url)"
# rather than dropped -- a real Word hyperlink field would need hand-built relationship
# XML, and losing the URL entirely is worse than showing it.
_INLINE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|__(?P<bold2>.+?)__"
    r"|(?<!\*)\*(?P<italic>[^*\n]+?)\*(?!\*)"
    r"|(?<![\w_])_(?P<italic2>[^_\n]+?)_(?![\w_])"
    r"|`(?P<code>[^`\n]+?)`"
    r"|\[(?P<link>[^\]\n]*)\]\((?P<url>[^)\s]*)[^)]*\)"
)

_MONOSPACE = "Consolas"


def _styled(document, style_name: str, text: str = ""):
    """A paragraph in `style_name`, falling back to the default style if this Word
    template doesn't define it. python-docx raises KeyError for an unknown style, and a
    missing style is never worth failing a student's download over."""
    try:
        return document.add_paragraph(text, style=style_name)
    except KeyError:
        return document.add_paragraph(text)


def _add_inline_runs(paragraph, text: str) -> None:
    """Splits one line of Markdown into real Word runs: bold, italic, inline code, and
    flattened links. Unmatched text becomes plain runs, so anything this doesn't
    understand still arrives intact rather than being swallowed."""
    cursor = 0
    for match in _INLINE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        groups = match.groupdict()
        if groups["bold"] is not None or groups["bold2"] is not None:
            paragraph.add_run(groups["bold"] or groups["bold2"]).bold = True
        elif groups["italic"] is not None or groups["italic2"] is not None:
            paragraph.add_run(groups["italic"] or groups["italic2"]).italic = True
        elif groups["code"] is not None:
            run = paragraph.add_run(groups["code"])
            run.font.name = _MONOSPACE
        else:
            label = groups["link"] or ""
            url = groups["url"] or ""
            paragraph.add_run(f"{label} ({url})" if url and url != label else label or url)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _flush_paragraph(document, buffer: list[str]) -> None:
    """Consecutive non-blank lines are one paragraph, matching how the student's own
    Preview tab renders them: a single newline is a Markdown soft break, and remark-gfm
    (which is what <MessageContent> uses) does not turn it into a hard line break."""
    if not buffer:
        return
    paragraph = document.add_paragraph()
    _add_inline_runs(paragraph, " ".join(buffer))
    buffer.clear()


def _add_table(document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    try:
        table.style = "Table Grid"  # borders; the default style draws none at all
    except KeyError:
        pass
    for row_index, cells in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            paragraph = cell.paragraphs[0]
            _add_inline_runs(paragraph, cells[column_index] if column_index < len(cells) else "")
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def build_document_docx(title: str, content: str) -> bytes:
    """Renders a document's title and Markdown content into real `.docx` bytes.

    Empty content is not an error: a brand-new note is a real document the student can
    legitimately want a (nearly empty, correctly titled) Word file for.
    """
    document = docx.Document()
    document.add_heading(title or "Untitled", level=0)

    lines = (content or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    buffer: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]

        if _FENCE.match(line):
            _flush_paragraph(document, buffer)
            index += 1
            code: list[str] = []
            while index < len(lines) and not _FENCE.match(lines[index]):
                code.append(lines[index])
                index += 1
            index += 1  # the closing fence (or the end of the document)
            for code_line in code:
                paragraph = _styled(document, "No Spacing")
                run = paragraph.add_run(code_line)
                run.font.name = _MONOSPACE
                run.font.size = Pt(10)
            continue

        if _TABLE_ROW.match(line) and index + 1 < len(lines) and _TABLE_DIVIDER.match(lines[index + 1]):
            _flush_paragraph(document, buffer)
            rows = [_split_table_row(line)]
            index += 2  # header + divider
            while index < len(lines) and _TABLE_ROW.match(lines[index]):
                rows.append(_split_table_row(lines[index]))
                index += 1
            _add_table(document, rows)
            continue

        index += 1

        if not line.strip():
            _flush_paragraph(document, buffer)
            continue

        if _RULE.match(line):
            _flush_paragraph(document, buffer)
            continue

        heading = _HEADING.match(line)
        if heading:
            _flush_paragraph(document, buffer)
            paragraph = _styled(document, f"Heading {len(heading.group(1))}")
            _add_inline_runs(paragraph, heading.group(2).strip())
            continue

        quote = _QUOTE.match(line)
        if quote:
            _flush_paragraph(document, buffer)
            _add_inline_runs(_styled(document, "Quote"), quote.group(1).strip())
            continue

        bullet = _BULLET.match(line)
        if bullet:
            _flush_paragraph(document, buffer)
            _add_inline_runs(_styled(document, "List Bullet"), bullet.group(1).strip())
            continue

        numbered = _NUMBERED.match(line)
        if numbered:
            _flush_paragraph(document, buffer)
            _add_inline_runs(_styled(document, "List Number"), numbered.group(1).strip())
            continue

        buffer.append(line.strip())

    _flush_paragraph(document, buffer)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
