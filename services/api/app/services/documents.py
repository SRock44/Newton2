import asyncio
import io
import uuid
from typing import Any

import docx
import pptx
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
from sqlalchemy import delete, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.object_storage import (
    ensure_bucket_sync,
    get_object_sync,
    put_object_sync,
    remove_object_sync,
)
from app.db.models import Document, DocumentChunk
from app.memory.rag import store_document_chunks

MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20MB

_TEXT_MIME_TYPES = {"text/plain", "text/markdown"}
_TEXT_EXTENSIONS = (".txt", ".md", ".markdown")
_PDF_EXTENSION = ".pdf"
# Office Open XML -- lecture slides and essays/handouts, by far the two most common
# things a student is actually handed after a PDF. Both are really zip archives of XML,
# so there's no "decode it as text" fallback: they need their own real parser
# (python-pptx / python-docx) exactly like a PDF needs pypdf.
_PPTX_EXTENSION = ".pptx"
_PPTX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_DOCX_EXTENSION = ".docx"
_DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUPPORTED_TYPES_MESSAGE = "Unsupported file type — upload .txt, .md, .pdf, .pptx, or .docx"


def _extract_pdf_text(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_pptx_text(raw: bytes) -> str:
    """Every text-bearing shape on every slide, in slide order then in the slide's own
    shape order -- which is what a student means by "what's on the slides", including
    titles, bullet bodies, and text inside tables. Grouped shapes are walked recursively
    (python-pptx exposes a group's members only through its own .shapes), and notes are
    deliberately left out: a lecturer's speaker notes are frequently absent, and when
    present they're a different kind of content from the deck itself.

    One blank line between slides so the RAG chunker (app/memory/rag.py) has a natural
    paragraph boundary to split on rather than running two slides together."""
    presentation = pptx.Presentation(io.BytesIO(raw))
    slides: list[str] = []
    for slide in presentation.slides:
        lines: list[str] = []
        _collect_pptx_shape_text(slide.shapes, lines)
        if lines:
            slides.append("\n".join(lines))
    return "\n\n".join(slides)


def _collect_pptx_shape_text(shapes: Any, lines: list[str]) -> None:
    for shape in shapes:
        # MSO_SHAPE_TYPE.GROUP == 6; compared numerically to avoid importing the enum
        # just for one check (python-pptx's own shape_type can also be None).
        if getattr(shape, "shape_type", None) == 6:
            _collect_pptx_shape_text(shape.shapes, lines)
            continue
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    lines.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                joined = " | ".join(c for c in cells if c)
                if joined:
                    lines.append(joined)


def _extract_docx_text(raw: bytes) -> str:
    """Every paragraph's text in document order, plus table cell rows (a lot of real
    coursework -- rubrics, problem sets, lab data -- lives in tables, and dropping it
    would silently lose content a student can plainly see in Word). Empty paragraphs are
    skipped rather than emitted as blank lines."""
    document = docx.Document(io.BytesIO(raw))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            joined = " | ".join(c for c in cells if c)
            if joined:
                lines.append(joined)
    return "\n".join(lines)


def _extract_text(filename: str, mime_type: str | None, raw: bytes) -> str:
    """Runs in a worker thread (see upload_document) since PDF/Office parsing is
    CPU-bound and would otherwise block the event loop."""
    lower_name = filename.lower()
    if lower_name.endswith(_PDF_EXTENSION) or mime_type == "application/pdf":
        return _extract_pdf_text(raw)
    if lower_name.endswith(_PPTX_EXTENSION) or mime_type == _PPTX_MIME_TYPE:
        try:
            return _extract_pptx_text(raw)
        except Exception as exc:  # noqa: BLE001 - a corrupt/mislabeled .pptx is a 400, not a 500
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST, "Couldn't read this PowerPoint file — it may be corrupt"
            ) from exc
    if lower_name.endswith(_DOCX_EXTENSION) or mime_type == _DOCX_MIME_TYPE:
        try:
            return _extract_docx_text(raw)
        except Exception as exc:  # noqa: BLE001 - same reasoning as .pptx above
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST, "Couldn't read this Word file — it may be corrupt"
            ) from exc
    if lower_name.endswith(_TEXT_EXTENSIONS) or mime_type in _TEXT_MIME_TYPES:
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "File is not valid UTF-8 text") from exc
    raise HTTPException(status.HTTP_400_BAD_REQUEST, SUPPORTED_TYPES_MESSAGE)


async def _store_document(
    db: AsyncSession,
    user_id: uuid.UUID,
    filename: str,
    mime_type: str | None,
    raw: bytes,
    text: str,
    kind: str = "upload",
    paper_sources: list[dict] | None = None,
) -> Document:
    """Shared tail end of upload_document(), upload_document_bytes(), and create_note():
    store the raw bytes in MinIO under a per-user key, create the Document row, then run
    the chunk+embed+store pipeline (app.memory.rag) over the already-extracted text.
    `kind` distinguishes a student upload from a Notepad note (see Document.kind's
    docstring in app/db/models.py) -- everything else about the pipeline is identical.
    `paper_sources` is the de-duplicated, final-keyed bibliography source list that
    produced this document, for a write_research_paper output only (see
    Document.paper_sources) -- None/[] for every other document, which is what makes
    GET /documents/{id}/bibliography.bib 404 for a plain upload."""
    settings = get_settings()
    document_id = uuid.uuid4()
    minio_key = f"{user_id}/{document_id}/{filename}"

    # MinIO's SDK is synchronous; run its blocking network calls off the event loop.
    await asyncio.to_thread(ensure_bucket_sync, settings.minio_bucket)
    await asyncio.to_thread(put_object_sync, settings.minio_bucket, minio_key, raw, mime_type)

    document = Document(
        id=document_id,
        user_id=user_id,
        filename=filename,
        mime_type=mime_type,
        minio_key=minio_key,
        kind=kind,
        paper_sources=paper_sources or None,
    )
    db.add(document)
    await db.flush()

    await store_document_chunks(db, document.id, text)
    await db.commit()
    return document


async def upload_document(db: AsyncSession, user_id: uuid.UUID, file: UploadFile) -> Document:
    """Validate the upload, store the raw bytes in MinIO under a per-user key, create the
    Document row, then run the chunk+embed+store pipeline (app.memory.rag) over the
    extracted text."""
    raw = await file.read()
    if not raw:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File is empty")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"File too large — max {MAX_UPLOAD_BYTES // (1024 * 1024)}MB",
        )

    filename = file.filename or "upload"
    text = await asyncio.to_thread(_extract_text, filename, file.content_type, raw)
    return await _store_document(db, user_id, filename, file.content_type, raw, text)


async def upload_document_bytes(
    db: AsyncSession,
    user_id: uuid.UUID,
    filename: str,
    mime_type: str,
    raw: bytes,
    paper_sources: list[dict] | None = None,
) -> Document:
    """Bytes-based sibling of upload_document() for content that never arrived as an
    HTTP UploadFile -- e.g. app/tools/write_research_paper.py's compiled PDF and
    generated .tex source, which exist only as in-memory bytes this process produced
    itself. Runs the exact same _extract_text() dispatch upload_document() uses (by
    filename/mime_type), so a generated PDF here is pypdf-extracted for RAG exactly like
    a student-uploaded PDF would be, and a generated .tex file (mime_type="text/plain")
    is stored as plain decoded text -- which for a .tex source IS the document's own
    content, no separate extraction step needed. No MAX_UPLOAD_BYTES check here: this
    content wasn't submitted by an HTTP client and is already bounded well under that
    ceiling by sandbox-runner's own LATEX_MAX_PDF_BYTES."""
    text = await asyncio.to_thread(_extract_text, filename, mime_type, raw)
    return await _store_document(db, user_id, filename, mime_type, raw, text, paper_sources=paper_sources)


async def create_note(db: AsyncSession, user_id: uuid.UUID, title: str) -> Document:
    """Creates a brand-new, empty note -- a first-class named Document with
    kind="note" (see app/routers/notes.py's "Newton Notepad" feature). Goes through
    the exact same MinIO-store + chunk+embed pipeline _store_document already uses for
    an upload; chunk_text("") yields no chunks, so this starts with zero DocumentChunk
    rows until the student's first real PATCH /notes/{id} save -- no separate code path
    to keep in sync with the real editing pipeline."""
    return await _store_document(db, user_id, title, "text/markdown", b"", "", kind="note")


async def get_document_text(document: Document) -> str:
    """Re-fetches the raw file from MinIO and re-runs extraction, rather than storing
    the full text separately from the (overlapping, chunked) RAG copy — documents here
    are small enough that re-extracting on demand is cheap, and it avoids keeping two
    representations of the same content in sync."""
    settings = get_settings()
    raw = await asyncio.to_thread(get_object_sync, settings.minio_bucket, document.minio_key)
    return await asyncio.to_thread(_extract_text, document.filename, document.mime_type, raw)


async def get_document_raw(document: Document) -> bytes:
    """Returns the exact bytes stored in MinIO for this document, with no text
    extraction — used to serve the raw file (e.g. embedding a PDF viewer)."""
    settings = get_settings()
    return await asyncio.to_thread(get_object_sync, settings.minio_bucket, document.minio_key)


def is_editable(document: Document) -> bool:
    """True only for plain text/markdown documents — PDFs are view-only since editing
    extracted PDF text and writing it back as a PDF isn't attempted (see the module
    docstring-level rationale in the router)."""
    lower_name = document.filename.lower()
    if lower_name.endswith(_PDF_EXTENSION) or document.mime_type == "application/pdf":
        return False
    return lower_name.endswith(_TEXT_EXTENSIONS) or document.mime_type in _TEXT_MIME_TYPES


async def update_document_content(
    db: AsyncSession, document: Document, content: str, filename: str | None = None
) -> Document:
    """Overwrites the stored file with edited text and re-chunks it for RAG, so
    retrieval always reflects exactly what the student sees and last edited. `filename`
    is optional -- PATCH /notes/{id} passes the note's (possibly renamed) title through
    here so a title-and-content save is one write, one re-chunk, one commit; the plain
    document-editor's PUT /documents/{id}/content never passes it and only touches
    content, exactly as before. Re-chunking the WHOLE document on every save (rather
    than diffing) is intentional -- these are small text documents/notes, not large
    PDFs, so it's not worth the complexity of incremental re-chunking."""
    settings = get_settings()
    raw = content.encode("utf-8")
    await asyncio.to_thread(
        put_object_sync, settings.minio_bucket, document.minio_key, raw, document.mime_type
    )
    await db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id))
    await store_document_chunks(db, document.id, content)
    if filename is not None:
        document.filename = filename
    # Explicitly touched (not just relying on onupdate=func.now()) because the
    # DocumentChunk delete+insert above doesn't itself dirty this Document row -- with
    # no filename change, this UPDATE would otherwise never be emitted at all.
    document.updated_at = func.now()
    await db.commit()
    await db.refresh(document)
    return document


async def delete_document(db: AsyncSession, document: Document) -> None:
    """Removes a document's chunks and DB row, then its raw object in MinIO."""
    settings = get_settings()
    await db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id))
    await db.delete(document)
    await db.commit()
    try:
        await asyncio.to_thread(remove_object_sync, settings.minio_bucket, document.minio_key)
    except Exception:
        # The DB rows are the source of truth for what the user sees; don't fail the
        # request over a MinIO object that's already gone or unreachable.
        pass
