"""Real tests for the two new "get what you made back out of Newton" document paths:

  1. .pptx / .docx ingestion (app/services/documents.py's _extract_text) -- exercised
     against REAL files generated here by python-pptx/python-docx themselves, not
     fixtures checked in or a mocked library call.
  2. GET /documents/{id}/bibliography.bib -- the .bib a research paper's own sources
     assemble into.

Endpoint behavior is exercised by calling the router coroutine directly with a synthetic
claims dict, exactly the way the real request would reach it after auth, so every test
runs against its OWN throwaway User (get_or_create_user JIT-provisions one from whatever
claims it's handed) and never touches the shared dev account's documents. Follows
tests/test_gamification.py's throwaway-user fixture + explicit-teardown convention.
"""

import io
import uuid

import pytest
import pytest_asyncio
from fastapi import HTTPException
from pypdf import PdfWriter
from sqlalchemy import delete, select

import docx
import pptx
from app.db.models import Document, DocumentChunk, User
from app.routers.documents import get_bibliography, list_documents
from app.services import documents as documents_service
from app.services.documents import _extract_text, upload_document_bytes

# ---------------------------------------------------------------------------
# Real .pptx / .docx files, built here with the same libraries that parse them.
# ---------------------------------------------------------------------------


def _make_pptx(slides: list[list[str]]) -> bytes:
    """One slide per inner list, one text box per string on it. Layout 6 is the blank
    layout, so nothing but the text actually added here ends up in the deck -- a title
    placeholder's default "Click to add title" prompt text would otherwise pollute the
    extraction and make this test lie about what it proved."""
    presentation = pptx.Presentation()
    blank = presentation.slide_layouts[6]
    for lines in slides:
        slide = presentation.slides.add_slide(blank)
        for index, line in enumerate(lines):
            box = slide.shapes.add_textbox(
                pptx.util.Inches(1), pptx.util.Inches(1 + index), pptx.util.Inches(6), pptx.util.Inches(1)
            )
            box.text_frame.text = line
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _make_docx(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_extract_text_reads_every_slide_of_a_real_pptx_in_order():
    raw = _make_pptx(
        [
            ["Photosynthesis", "Light-dependent reactions happen in the thylakoid."],
            ["The Calvin cycle fixes carbon in the stroma."],
            ["ATP and NADPH power the cycle."],
        ]
    )

    text = _extract_text("lecture-4.pptx", PPTX_MIME, raw)

    assert "Photosynthesis" in text
    assert "thylakoid" in text
    assert "Calvin cycle" in text
    assert "NADPH" in text
    # Slide order preserved, not a set-like jumble.
    assert text.index("Photosynthesis") < text.index("Calvin cycle") < text.index("NADPH")
    # Slides separated by a blank line so the RAG chunker has a real paragraph boundary.
    assert "\n\n" in text


def test_extract_text_reads_a_real_pptx_recognized_by_extension_alone():
    """A browser/OS that hands up a generic content type (or none) must still work --
    the extension is a first-class signal here, same as it already is for .pdf."""
    raw = _make_pptx([["Slide one body text"]])
    assert "Slide one body text" in _extract_text("deck.pptx", "application/octet-stream", raw)


def test_extract_text_reads_pptx_table_cells():
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(
        2, 2, pptx.util.Inches(1), pptx.util.Inches(1), pptx.util.Inches(6), pptx.util.Inches(2)
    )
    shape.table.cell(0, 0).text = "Element"
    shape.table.cell(0, 1).text = "Symbol"
    shape.table.cell(1, 0).text = "Sodium"
    shape.table.cell(1, 1).text = "Na"
    buffer = io.BytesIO()
    presentation.save(buffer)

    text = _extract_text("periodic.pptx", PPTX_MIME, buffer.getvalue())

    assert "Sodium" in text and "Na" in text


def test_extract_text_reads_every_paragraph_of_a_real_docx_in_order():
    raw = _make_docx(
        [
            "Essay: The Causes of the French Revolution",
            "",
            "Fiscal crisis was the proximate trigger.",
            "The Estates-General convened in May 1789.",
        ]
    )

    text = _extract_text("essay.docx", DOCX_MIME, raw)

    assert "French Revolution" in text
    assert "Fiscal crisis" in text
    assert "Estates-General" in text
    assert text.index("Fiscal crisis") < text.index("Estates-General")
    # Empty paragraphs are skipped rather than emitted as blank lines.
    assert "\n\n" not in text


def test_extract_text_reads_a_real_docx_recognized_by_extension_alone():
    raw = _make_docx(["Only paragraph."])
    assert "Only paragraph." in _extract_text("paper.docx", None, raw)


def test_extract_text_reads_docx_table_cells():
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Assignment"
    table.cell(0, 1).text = "Due"
    table.cell(1, 0).text = "Problem set 3"
    table.cell(1, 1).text = "October 14"
    buffer = io.BytesIO()
    document.save(buffer)

    text = _extract_text("rubric.docx", DOCX_MIME, buffer.getvalue())

    assert "Problem set 3" in text and "October 14" in text


def test_extract_text_rejects_a_corrupt_pptx_as_a_400_not_a_500():
    with pytest.raises(HTTPException) as excinfo:
        _extract_text("broken.pptx", PPTX_MIME, b"this is definitely not a zip archive")
    assert excinfo.value.status_code == 400


def test_extract_text_rejects_a_corrupt_docx_as_a_400_not_a_500():
    with pytest.raises(HTTPException) as excinfo:
        _extract_text("broken.docx", DOCX_MIME, b"nope")
    assert excinfo.value.status_code == 400


def test_extract_text_still_rejects_a_genuinely_unsupported_type():
    with pytest.raises(HTTPException) as excinfo:
        _extract_text("archive.zip", "application/zip", b"PK\x03\x04junk")
    assert excinfo.value.status_code == 400
    # The message now names the two new types it really does accept.
    assert ".pptx" in excinfo.value.detail and ".docx" in excinfo.value.detail


def test_office_documents_are_not_editable():
    """PDFs are view-only because round-tripping extracted text back into a PDF would
    produce garbage; a .pptx/.docx is the same situation, so is_editable must say no."""
    assert not documents_service.is_editable(Document(filename="deck.pptx", mime_type=PPTX_MIME))
    assert not documents_service.is_editable(Document(filename="essay.docx", mime_type=DOCX_MIME))
    assert documents_service.is_editable(Document(filename="notes.md", mime_type="text/markdown"))


# ---------------------------------------------------------------------------
# End-to-end ingestion: a real .pptx becomes a real Document with real RAG chunks.
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture
async def throwaway_user(db_session):
    user = User(keycloak_sub=f"test-document-export-{uuid.uuid4()}")
    db_session.add(user)
    await db_session.flush()
    await db_session.commit()

    yield user

    document_ids = (
        (await db_session.execute(select(Document.id).where(Document.user_id == user.id)))
        .scalars()
        .all()
    )
    if document_ids:
        await db_session.execute(delete(DocumentChunk).where(DocumentChunk.document_id.in_(document_ids)))
    await db_session.execute(delete(Document).where(Document.user_id == user.id))
    await db_session.execute(delete(User).where(User.id == user.id))
    await db_session.commit()


def _claims(user: User) -> dict:
    return {"sub": user.keycloak_sub}


@pytest.mark.asyncio
async def test_a_real_pptx_ingests_into_a_document_with_searchable_chunks(db_session, throwaway_user):
    raw = _make_pptx([["Mitosis", "Prophase, metaphase, anaphase, telophase."]])

    document = await upload_document_bytes(
        db_session, throwaway_user.id, "cell-division.pptx", PPTX_MIME, raw
    )

    chunks = (
        (await db_session.execute(select(DocumentChunk).where(DocumentChunk.document_id == document.id)))
        .scalars()
        .all()
    )
    assert chunks, "a .pptx upload must produce real RAG chunks, not just a stored blob"
    assert any("metaphase" in c.content for c in chunks)
    # And the stored bytes round-trip unchanged (what a Download would hand back).
    assert await documents_service.get_document_raw(document) == raw


@pytest.mark.asyncio
async def test_a_real_docx_ingests_into_a_document_with_searchable_chunks(db_session, throwaway_user):
    raw = _make_docx(["Thesis statement about the Treaty of Westphalia and sovereignty."])

    document = await upload_document_bytes(
        db_session, throwaway_user.id, "history-essay.docx", DOCX_MIME, raw
    )

    chunks = (
        (await db_session.execute(select(DocumentChunk).where(DocumentChunk.document_id == document.id)))
        .scalars()
        .all()
    )
    assert any("Westphalia" in c.content for c in chunks)


# ---------------------------------------------------------------------------
# GET /documents/{id}/bibliography.bib
# ---------------------------------------------------------------------------

PAPER_SOURCES = [
    {
        "key": "doe2024attention",
        "type": "article",
        "author": "Jane Doe",
        "title": "Attention Is All You Need",
        "year": "2024",
        "venue": "Journal of Irreproducible Results",
        "url": "https://example.org/attention",
    },
    {
        "key": "smith2020systems",
        "type": "inproceedings",
        "author": "Smith, Alan",
        "title": "Analysis of Systems",
        "year": "2020",
        "venue": "Proceedings of Everything",
    },
]


@pytest.mark.asyncio
async def test_bibliography_endpoint_returns_a_real_bib_for_a_paper_with_sources(db_session, throwaway_user):
    document = await upload_document_bytes(
        db_session,
        throwaway_user.id,
        "My Paper.pdf",
        "application/pdf",
        _minimal_pdf(),
        paper_sources=PAPER_SOURCES,
    )

    response = await get_bibliography(document.id, claims=_claims(throwaway_user), db=db_session)

    body = response.body.decode("utf-8")
    assert response.status_code == 200
    assert "application/x-bibtex" in response.headers["content-type"]
    # Real BibTeX entries with the SAME keys the paper's own \cite{} placeholders use.
    assert "@article{doe2024attention," in body
    assert "@inproceedings{smith2020systems," in body
    assert "journal = {Journal of Irreproducible Results}" in body
    assert "booktitle = {Proceedings of Everything}" in body
    assert "url = {https://example.org/attention}" in body
    # Named after the paper, so it lands next to it in a Downloads folder.
    assert 'filename="My Paper.bib"' in response.headers["content-disposition"]


@pytest.mark.asyncio
async def test_bibliography_endpoint_404s_for_a_plain_upload(db_session, throwaway_user):
    document = await upload_document_bytes(
        db_session, throwaway_user.id, "syllabus.txt", "text/plain", b"Week 1: intro"
    )

    with pytest.raises(HTTPException) as excinfo:
        await get_bibliography(document.id, claims=_claims(throwaway_user), db=db_session)
    assert excinfo.value.status_code == 404


@pytest.mark.asyncio
async def test_bibliography_endpoint_404s_for_someone_elses_document(db_session, throwaway_user):
    other = User(keycloak_sub=f"test-document-export-other-{uuid.uuid4()}")
    db_session.add(other)
    await db_session.flush()
    document = await upload_document_bytes(
        db_session, other.id, "Their Paper.pdf", "application/pdf", _minimal_pdf(), paper_sources=PAPER_SOURCES
    )
    try:
        with pytest.raises(HTTPException) as excinfo:
            await get_bibliography(document.id, claims=_claims(throwaway_user), db=db_session)
        assert excinfo.value.status_code == 404
    finally:
        await db_session.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id))
        await db_session.execute(delete(Document).where(Document.user_id == other.id))
        await db_session.execute(delete(User).where(User.id == other.id))
        await db_session.commit()


@pytest.mark.asyncio
async def test_document_listing_flags_which_documents_have_a_bibliography(db_session, throwaway_user):
    await upload_document_bytes(
        db_session,
        throwaway_user.id,
        "Paper.pdf",
        "application/pdf",
        _minimal_pdf(),
        paper_sources=PAPER_SOURCES,
    )
    await upload_document_bytes(db_session, throwaway_user.id, "notes.txt", "text/plain", b"hello")

    rows = await list_documents(claims=_claims(throwaway_user), db=db_session)

    by_name = {row["filename"]: row for row in rows}
    assert by_name["Paper.pdf"]["has_bibliography"] is True
    assert by_name["notes.txt"]["has_bibliography"] is False


def _minimal_pdf() -> bytes:
    """A minimal, REAL, parseable single-page PDF -- not just bytes with a .pdf name --
    so upload_document_bytes' own pypdf extraction succeeds. Same helper shape as
    tests/test_write_research_paper_tool.py's `_minimal_real_pdf_bytes`."""
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)
    return buffer.getvalue()
